"""Tests for :mod:`siscadro_rw5.extractor` and the package-level API."""

from __future__ import annotations

import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from siscadro_survey.records import IssueSeverity, SourceFormat

from siscadro_rw5 import export_to_database, export_to_xlsx, extract_points
from siscadro_rw5.extractor import Rw5Extractor


class TestCanRead:
    """Tests for :meth:`Rw5Extractor.can_read`."""

    def test_accepts_rw5_extension(self):
        assert Rw5Extractor().can_read(Path("job.rw5")) is True

    def test_accepts_rw5_extension_case_insensitively(self):
        assert Rw5Extractor().can_read(Path("JOB.RW5")) is True

    def test_rejects_other_extensions(self):
        assert Rw5Extractor().can_read(Path("job.txt")) is False

    def test_declares_its_format_name_and_extensions(self):
        extractor = Rw5Extractor()
        assert extractor.format_name == "rw5"
        assert extractor.extensions == (".rw5",)


class TestExtract:
    """Tests for :meth:`Rw5Extractor.extract`."""

    def test_source_metadata_describes_the_file(
        self, english_minimal_path: Path
    ):
        result = Rw5Extractor().extract(english_minimal_path)

        assert result.source.source_type == SourceFormat.RW5
        assert result.source.resolved_path == english_minimal_path.resolve()
        assert result.source.original_name == "english_minimal.rw5"
        assert result.source.parser_metadata["job_name"] == "JOB-EN-01"
        assert result.source.parser_metadata["locale"] == "en"
        assert result.source.parser_metadata["equipment"] == (
            "Robotic Total Station"
        )

    def test_extracts_two_canonical_records(self, english_minimal_path: Path):
        result = Rw5Extractor().extract(english_minimal_path)

        assert len(result.records) == 2

    def test_maps_coordinates_as_decimals(self, english_minimal_path: Path):
        result = Rw5Extractor().extract(english_minimal_path)

        first = result.records[0]
        assert isinstance(first.north, Decimal)
        assert first.north == pytest.approx(Decimal("500000.123"))
        assert first.east == pytest.approx(Decimal("300000.456"))
        assert first.height == pytest.approx(Decimal("120.789"))

    def test_maps_identity_geographic_and_quality_fields(
        self, english_minimal_path: Path
    ):
        result = Rw5Extractor().extract(english_minimal_path)

        first = result.records[0]
        assert first.name == "101"
        assert first.code == "CP1"
        assert first.latitude == pytest.approx(46.500100000)
        assert first.longitude == pytest.approx(24.500200000)
        assert first.wgs84_altitude == pytest.approx(500.5)
        assert first.observed_at_utc == datetime.datetime(
            2026, 7, 15, 9, 5, 12, tzinfo=datetime.timezone.utc
        )
        assert first.method == "Average"
        assert first.status == "FIXED"
        assert first.satellite_count == 12
        assert first.hrms == pytest.approx(0.015)
        assert first.vrms == pytest.approx(0.025)
        assert first.hdop == pytest.approx(1.2)
        assert first.vdop == pytest.approx(1.4)
        assert first.pdop == pytest.approx(1.8)
        assert first.tdop == pytest.approx(1.1)
        assert first.gdop == pytest.approx(2.2)

    def test_maps_antenna_and_base_fields(self, english_minimal_path: Path):
        result = Rw5Extractor().extract(english_minimal_path)

        first = result.records[0]
        assert first.antenna_measurement_method == "Vertical"
        assert first.entered_antenna_height == pytest.approx(1.8)
        assert first.true_antenna_height == pytest.approx(1.8)
        assert first.base_id == "1"
        assert first.base_latitude == pytest.approx(46.5)
        assert first.base_longitude == pytest.approx(24.5)
        assert first.base_height == pytest.approx(500.0)
        assert first.source_record_id == "101"

    def test_collects_unmapped_fields_into_source_values(
        self, english_minimal_path: Path
    ):
        result = Rw5Extractor().extract(english_minimal_path)

        first = result.records[0]
        assert first.source_values["comment"] == "CP1"
        assert "north_avg" not in first.source_values

    def test_reports_a_parse_issue_for_the_malformed_line(
        self, english_minimal_path: Path
    ):
        result = Rw5Extractor().extract(english_minimal_path)

        assert any(
            issue.severity == IssueSeverity.WARNING for issue in result.issues
        )

    def test_includes_stakeout_data_in_source_values(
        self, quality_and_offsets_path: Path
    ):
        result = Rw5Extractor().extract(quality_and_offsets_path)

        stakeout_record = next(
            record for record in result.records if record.name == "301"
        )
        assert stakeout_record.source_values["stakeout"]["Design Pt#"] == (
            "2061"
        )
        assert stakeout_record.source_values["stakeout_source"] == (
            "X DESIGN.crd"
        )

    def test_skips_a_point_missing_projected_coordinates(self, tmp_path: Path):
        path = tmp_path / "missing_gs.rw5"
        path.write_text(
            "\n".join(
                [
                    "--Equipment: Test",
                    "BP,PN1,LA1,LN1,EL1,AG0,PA0",
                    "GPS,LA1,LN1,EL1,PN101,--C",
                    "G0,07/17/2026 11:00:00,Base ID read at rover: 1",
                ]
            ),
            encoding="utf-8",
        )

        result = Rw5Extractor().extract(path)

        assert result.records == ()
        assert any(
            issue.severity == IssueSeverity.WARNING
            and "missing" in issue.message
            and issue.record_id == "101"
            for issue in result.issues
        )

    def test_silently_skips_vrs_network_base_gps(self, tmp_path: Path):
        path = tmp_path / "vrs_base.rw5"
        path.write_text(
            "\n".join(
                [
                    "--Equipment: S9 III GNSS",
                    "--RTK Method: RTCM V3.0,Device: Internal GSM,"
                    "Network: NTRIP RO_VRS_3.1_GG",
                    "JB,NMCompl2,DT09-29-2016,TM12:45:24",
                    "BP,PN461,LA45.75981728500,LN25.15032520194,EL500.000,"
                    "AG0.000,PA0.088",
                    "GPS,PNRO_VRS_3.1_GG,LA45.75981728500,LN25.15032520194,"
                    "EL500.000,--",
                    "G0,09/29/2016 12:40:00,Base ID read at rover: 0460",
                    "GPS,PN1,LA45.74096737583,LN25.09619256111,EL579.903,"
                    "--POD",
                    "GS,PN1,N 471251.4086,E 507604.6572,EL537.9049,--POD",
                ]
            ),
            encoding="utf-8",
        )

        result = Rw5Extractor().extract(path)

        assert [record.name for record in result.records] == ["1"]
        assert not any(
            issue.record_id == "RO_VRS_3.1_GG" for issue in result.issues
        )
        assert not any(
            "missing projected" in issue.message for issue in result.issues
        )

    def test_silently_skips_base_number_gps_without_gs(self, tmp_path: Path):
        path = tmp_path / "base_pn_echo.rw5"
        path.write_text(
            "\n".join(
                [
                    "--Equipment: Test",
                    "BP,PN1,LA1,LN1,EL1,AG0,PA0",
                    "GPS,LA1,LN1,EL1,PN1,--C",
                    "G0,07/17/2026 11:00:00,Base ID read at rover: 1",
                ]
            ),
            encoding="utf-8",
        )

        result = Rw5Extractor().extract(path)

        assert result.records == ()
        assert not any(
            "missing projected" in issue.message for issue in result.issues
        )

    def test_base_only_gt_does_not_warn_missing_projected(
        self, tmp_path: Path
    ):
        """Base ``GT`` must not invent a ``?`` stub that warns on extract."""
        path = tmp_path / "base_only_gt.rw5"
        path.write_text(
            "\n".join(
                [
                    "--Equipment: Stonex, S9III+",
                    "BP,PNBP001,LA44.381767044221,LN26.582728721507,"
                    "EL73.5710,AG1.9308,PA0.1068,ATAPC,SRBASE,--",
                    "--GS,PNBP001,N 350637.7785,E 656734.4090,EL37.7220,"
                    "--Base",
                    "--GT,PNBP001,SW-522,ST-258074000,EW-522,ET-258074000",
                ]
            ),
            encoding="utf-8",
        )

        result = Rw5Extractor().extract(path)

        assert [record.name for record in result.records] == ["BP001"]
        assert result.records[0].kind == "base"
        assert not any(
            "missing projected" in issue.message for issue in result.issues
        )
        assert not any(
            issue.record_id == "?" for issue in result.issues
        )


class TestBaseConfigurationExtract:
    """Canonical extraction of SurvCE base-configuration blocks."""

    def test_emits_kind_base_and_rover_without_issues(
        self, base_configuration_path: Path
    ):
        result = Rw5Extractor().extract(base_configuration_path)

        names = [record.name for record in result.records]
        assert "1508" in names
        assert "1498" in names
        assert "1" in names
        base_records = [r for r in result.records if r.kind == "base"]
        assert len(base_records) == 2
        base_1508 = next(r for r in base_records if r.name == "1508")
        assert base_1508.north == pytest.approx(Decimal("467880.2788"))
        assert base_1508.source_values["configured_by_gps_position"] is True
        rover = next(r for r in result.records if r.name == "1")
        assert rover.kind == "gps"
        assert not any(
            "missing projected" in issue.message for issue in result.issues
        )
        assert not any(
            "failed to parse" in issue.message for issue in result.issues
        )

    def test_silently_skips_pre_bp_stakeout_placeholder(self, tmp_path: Path):
        path = tmp_path / "pre_bp_stakeout.rw5"
        path.write_text(
            "\n".join(
                [
                    "--Equipment: Test",
                    "--CTSS1:Halchiu.crd",
                    "--CTSD0:Design Pt#,Design Elv",
                    "--CTSD0:51,0.000",
                    "--RTK Method: RTCM V3.0, Device: Internal GSM, "
                    "Network: NTRIP RO_VRS_MSM4",
                    "BP,PN053,LA45.5,LN25.4,EL580.0,AG0.000,PA0.088",
                    "GPS,PN100,LA45.4,LN25.2,EL567.0,--GM",
                    "GS,PN100,N 475210.3,E 538893.5,EL527.0,--GM",
                ]
            ),
            encoding="utf-8",
        )

        result = Rw5Extractor().extract(path)

        assert [record.name for record in result.records] == ["100"]
        assert not any(
            "missing projected" in issue.message for issue in result.issues
        )

    def test_observed_at_utc_falls_back_to_local_time_without_g0_or_gt(
        self, offset_local_time_path: Path
    ):
        result = Rw5Extractor().extract(offset_local_time_path)

        by_name = {record.name: record for record in result.records}
        assert by_name["135"].observed_at_utc == datetime.datetime(
            2016, 8, 8, 11, 13, 49, tzinfo=datetime.timezone.utc
        )
        assert by_name["135"].observed_at_local == datetime.datetime(
            2016, 8, 8, 11, 13, 49
        )
        assert by_name["136"].observed_at_utc == datetime.datetime(
            2016, 8, 8, 11, 14, 31, tzinfo=datetime.timezone.utc
        )

    def test_observed_at_utc_prefers_g0_moment_over_local_time(
        self, offset_g0_moment_path: Path
    ):
        result = Rw5Extractor().extract(offset_g0_moment_path)

        by_name = {record.name: record for record in result.records}
        pn143 = by_name["143"]
        assert pn143.observed_at_utc == datetime.datetime(
            2015, 7, 9, 6, 39, 16, tzinfo=datetime.timezone.utc
        )
        assert pn143.observed_at_local == datetime.datetime(
            2015, 7, 9, 9, 39, 2
        )
        assert pn143.source_values["moment"] == "2015-07-09T06:39:16"
        assert pn143.source_values["local_time"] == "2015-07-09T09:39:02"

    def test_observed_at_utc_still_prefers_gt_when_present(
        self, english_minimal_path: Path
    ):
        result = Rw5Extractor().extract(english_minimal_path)

        first = result.records[0]
        assert first.observed_at_utc == datetime.datetime(
            2026, 7, 15, 9, 5, 12, tzinfo=datetime.timezone.utc
        )
        assert first.observed_at_local == datetime.datetime(
            2026, 7, 15, 9, 5, 12
        )


class TestPackageLevelApi:
    """Tests for the ``extract_points``/``export_to_xlsx``/
    ``export_to_database`` convenience wrappers."""

    def test_extract_points_delegates_to_rw5_extractor(
        self, romanian_minimal_path: Path
    ):
        result = extract_points(romanian_minimal_path)

        assert len(result.records) == 2
        assert result.source.source_type == SourceFormat.RW5

    def test_export_to_xlsx_writes_a_points_workbook(
        self, english_minimal_path: Path, tmp_path: Path
    ):
        output_path = tmp_path / "out.xlsx"

        summary = export_to_xlsx(english_minimal_path, output_path)

        assert summary.written_row_count == 2
        assert output_path.exists()
        workbook = openpyxl.load_workbook(output_path)
        assert workbook.sheetnames == ["Points", "Source"]
        header_row = next(workbook["Points"].iter_rows(values_only=True))
        assert "North" in header_row
        assert "Name" in header_row

    def test_export_to_xlsx_refuses_to_overwrite_by_default(
        self, english_minimal_path: Path, tmp_path: Path
    ):
        output_path = tmp_path / "out.xlsx"
        output_path.write_text("existing")

        with pytest.raises(FileExistsError):
            export_to_xlsx(english_minimal_path, output_path)

    def test_export_to_database_imports_records_by_sqlite_path(
        self, quality_and_offsets_path: Path, tmp_path: Path
    ):
        db_path = tmp_path / "survey.sqlite"

        summary = export_to_database(quality_and_offsets_path, db_path)

        assert summary.source_count == 1
        assert summary.discovered_records == 2
        assert summary.unique_points_inserted == 2
        assert summary.existing_points_linked == 0

    def test_export_to_database_is_idempotent_for_the_same_file(
        self, quality_and_offsets_path: Path, tmp_path: Path
    ):
        db_path = tmp_path / "survey.sqlite"
        export_to_database(quality_and_offsets_path, db_path)

        summary = export_to_database(quality_and_offsets_path, db_path)

        assert summary.skipped_identical_sources == 1
        assert summary.unique_points_inserted == 0
